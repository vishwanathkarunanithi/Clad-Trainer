from flask import Flask, request, jsonify, send_from_directory
import openpyxl
from openpyxl import Workbook
import os

app = Flask(__name__, static_folder='.')

EXCEL_FILE = "Student_Results.xlsx"

@app.route('/')
def serve_index():
    return send_from_directory('.', 'index.html')

@app.route('/<path:path>')
def serve_static(path):
    return send_from_directory('.', path)

@app.route('/save-result', methods=['POST'])
def save_result():
    data = request.json
    if not data:
        return jsonify({"error": "No data provided"}), 400
        
    test_id = data.get("testId", "Unknown")
    sheet_name = f"Mock Test {test_id}"
    
    # Check if excel file exists
    if os.path.exists(EXCEL_FILE):
        wb = openpyxl.load_workbook(EXCEL_FILE)
    else:
        wb = Workbook()
        # Remove default sheet if we are creating a new one
        if "Sheet" in wb.sheetnames:
            del wb["Sheet"]
            
    # Check if sheet exists for this test
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
    else:
        ws = wb.create_sheet(title=sheet_name)
        # Create headers for new sheet
        ws.append(["Date", "Name", "Register Number", "Class & Section", "Score (%)", "Attempted", "Total Questions"])
        
    # Append the student data
    row = [
        data.get("date", ""),
        data.get("name", ""),
        data.get("reg", ""),
        data.get("classSec", ""),
        data.get("score", ""),
        data.get("attempted", ""),
        data.get("total", "")
    ]
    ws.append(row)
    
    # Save file
    try:
        wb.save(EXCEL_FILE)
        return jsonify({"success": True, "message": "Result saved successfully!"})
    except Exception as e:
        return jsonify({"error": str(e)}), 500

if __name__ == '__main__':
    print("Starting server on http://localhost:5000")
    print("Please use this URL to access the test instead of opening index.html directly!")
    app.run(host='0.0.0.0', port=5000, debug=True)
